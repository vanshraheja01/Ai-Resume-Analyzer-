"""Resume text extraction and structured field parsing.

This is deliberately *not* AI-based — it's classic regex/heuristic NLP,
resilient to different resume layouts but not perfect (a resume with an
unusual structure may end up with some empty fields). The AI analysis layer
(Phase 4) reasons over `raw_text` + this structured output; it doesn't
replace this step, it builds on it.
"""

import re
from pathlib import Path
from typing import Any

import docx
import pymupdf


class ResumeParsingError(Exception):
    """Raised when a file can't be opened/read at all (corrupted, unsupported, empty)."""


def extract_text(file_path: Path, file_type: str) -> str:
    if file_type == "pdf":
        return _extract_text_from_pdf(file_path)
    if file_type == "docx":
        return _extract_text_from_docx(file_path)
    raise ResumeParsingError(f"Unsupported file type: {file_type}")


def _extract_text_from_pdf(file_path: Path) -> str:
    try:
        with pymupdf.open(file_path) as document:
            text = "\n".join(page.get_text() for page in document)
    except Exception as exc:  # PyMuPDF raises its own RuntimeError/ValueError on corrupt files
        raise ResumeParsingError("Could not read this PDF - it may be corrupted or password-protected.") from exc

    if not text.strip():
        raise ResumeParsingError("No extractable text found in this PDF (it may be a scanned image).")
    return text


def _extract_text_from_docx(file_path: Path) -> str:
    try:
        document = docx.Document(file_path)
        paragraphs = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                paragraphs.extend(cell.text for cell in row.cells)
        text = "\n".join(paragraphs)
    except Exception as exc:
        raise ResumeParsingError("Could not read this DOCX file - it may be corrupted.") from exc

    if not text.strip():
        raise ResumeParsingError("No extractable text found in this document.")
    return text


# --- Structured field extraction -------------------------------------------------

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(?:\+?\d{1,3}[\s.-]?)?\(?\d{2,4}\)?[\s.-]?\d{3,4}[\s.-]?\d{3,6}")
URL_RE = re.compile(r"(https?://\S+|(?:www\.)?(?:linkedin|github)\.com/\S+)", re.IGNORECASE)

SECTION_HEADERS: dict[str, list[str]] = {
    "education": ["education", "academic background", "qualifications"],
    "skills": ["skills", "technical skills", "core competencies", "key skills"],
    "experience": ["experience", "work experience", "professional experience", "employment history"],
    "projects": ["projects", "personal projects", "academic projects"],
    "certifications": ["certifications", "certificates", "licenses"],
    "achievements": ["achievements", "awards", "honors", "accomplishments"],
    "languages": ["languages", "language proficiency"],
}


def _looks_like_header(line: str) -> str | None:
    normalized = line.strip().strip(":").lower()
    if not normalized or len(normalized.split()) > 5:
        return None
    for canonical, aliases in SECTION_HEADERS.items():
        if normalized in aliases:
            return canonical
    return None


def _split_into_sections(lines: list[str]) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = {key: [] for key in SECTION_HEADERS}
    current_section: str | None = None

    for line in lines:
        header = _looks_like_header(line)
        if header is not None:
            current_section = header
            continue
        stripped = line.strip(" \t•-·*")
        if current_section and stripped:
            sections[current_section].append(stripped)

    return sections


def _extract_name(lines: list[str]) -> str | None:
    # Heuristic: the first non-empty line near the top that isn't contact info
    # and isn't itself a section header. Resumes overwhelmingly lead with the name.
    for line in lines[:5]:
        candidate = line.strip()
        if not candidate:
            continue
        if EMAIL_RE.search(candidate) or PHONE_RE.search(candidate) or URL_RE.search(candidate):
            continue
        if _looks_like_header(candidate):
            continue
        if len(candidate.split()) <= 5:
            return candidate
    return None


def _extract_location(lines: list[str]) -> str | None:
    # Heuristic: a short "City, State"/"City, Country"-shaped line near the top.
    location_pattern = re.compile(r"^[A-Za-z .]+,\s*[A-Za-z .]+$")
    for line in lines[:8]:
        candidate = line.strip()
        if location_pattern.match(candidate) and len(candidate) < 60:
            return candidate
    return None


def _split_list_field(section_lines: list[str]) -> list[str]:
    items: list[str] = []
    for line in section_lines:
        parts = re.split(r"[,|]", line) if "," in line or "|" in line else [line]
        items.extend(p.strip() for p in parts if p.strip())
    # De-duplicate while preserving order.
    seen: set[str] = set()
    deduped = []
    for item in items:
        key = item.lower()
        if key not in seen:
            seen.add(key)
            deduped.append(item)
    return deduped


def parse_resume(raw_text: str) -> dict[str, Any]:
    lines = [line for line in raw_text.splitlines()]
    non_empty_lines = [line for line in lines if line.strip()]

    emails = EMAIL_RE.findall(raw_text)
    phones = PHONE_RE.findall(raw_text)
    links = URL_RE.findall(raw_text)

    sections = _split_into_sections(lines)

    return {
        "name": _extract_name(non_empty_lines),
        "email": emails[0] if emails else None,
        "phone": next((p for p in phones if len(re.sub(r"\D", "", p)) >= 7), None),
        "location": _extract_location(non_empty_lines),
        "links": list(dict.fromkeys(links)),  # de-duplicated, order preserved
        "education": sections["education"],
        "skills": _split_list_field(sections["skills"]),
        "experience": sections["experience"],
        "projects": sections["projects"],
        "certifications": sections["certifications"],
        "achievements": sections["achievements"],
        "languages": _split_list_field(sections["languages"]),
    }
